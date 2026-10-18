"""CSV・Excel（.xlsx・.xlsm）の計測テーブルを読み込む。

- CSV は UTF-8（BOM あり・なし）を先に試し、だめなら Shift_JIS（cp932。日本語版 Windows の Excel が書き出す形式）で読む。
- 値はいったん文字列のまま読み、数値・日時への変換は前処理で行う（Wafer ID の先頭の 0 などを失わないため）。
  空欄だけを欠損にする（pandas の既定では "NA"・"null" などの文字も欠損になってしまう）。
- Excel は openpyxl の読み取り専用モードで1行ずつ読み、進み具合を返す（大きなファイルでも画面が止まって見えないように）。
- 古い .xls は読めない（openpyxl が対応していない）。
"""
from __future__ import annotations

import csv
from pathlib import Path

import pandas as pd

from .errors import UserFacingError, check_cancel

CSV_SUFFIXES = (".csv", ".txt", ".tsv")
EXCEL_SUFFIXES = (".xlsx", ".xlsm")
SUPPORTED_SUFFIXES = CSV_SUFFIXES + EXCEL_SUFFIXES
CSV_ENCODINGS = ("utf-8-sig", "cp932")
SNIFF_BYTES = 65536
EXCEL_PROGRESS_ROWS = 5000


def file_kind(path: str | Path) -> str:
    suffix = Path(path).suffix.lower()
    if suffix in CSV_SUFFIXES:
        return "csv"
    if suffix in EXCEL_SUFFIXES:
        return "excel"
    if suffix == ".xls":
        raise UserFacingError(
            "古い形式の Excel ファイル（.xls）は読み込めません。",
            "このアプリが使う読み込み部品（openpyxl）が .xls に対応していないためです。",
            "Excel でファイルを開き、「名前を付けて保存」で .xlsx か CSV（UTF-8）にしてから読み込んでください。",
        )
    raise UserFacingError(
        f"このファイルの形式（{suffix or '拡張子なし'}）は読み込めません。",
        "読み込めるのは CSV（.csv・.txt・.tsv）と Excel（.xlsx・.xlsm）だけです。",
        "Excel で開いて .xlsx か CSV で保存し直してください。",
    )


def list_sheets(path: str | Path) -> list[str]:
    """Excel のシート名（CSV は空のリスト）。"""
    if file_kind(path) != "excel":
        return []
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:  # 壊れたファイル・パスワード付きなど
        raise _open_error(path, error) from error
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def read_table(path: str | Path, sheet: str | None = None, progress=None, cancel=None) -> pd.DataFrame:
    """計測テーブルを DataFrame で返す（列名は文字列、値は文字列・数値・日時のまま）。"""
    path = Path(path)
    if not path.exists():
        raise UserFacingError(
            "ファイルが見つかりません。",
            f"{path} が存在しないか、移動・削除された可能性があります。",
            "「ファイルを選ぶ…」でもう一度選び直してください。",
        )
    kind = file_kind(path)
    table = _read_csv(path) if kind == "csv" else _read_excel(path, sheet, progress, cancel)
    table = table.dropna(how="all")
    if table.empty:
        raise UserFacingError(
            "ファイルにデータの行がありません。",
            "1行目（見出し）より下に値が入っていないか、選んだシートが空です。",
            "ファイルの中身と、Excel の場合は選んだシートを確かめてください。",
        )
    return table.reset_index(drop=True)


def _read_csv(path: Path) -> pd.DataFrame:
    last_error = None
    for encoding in CSV_ENCODINGS:
        try:
            with open(path, encoding=encoding, newline="") as file:
                head = file.read(SNIFF_BYTES)
            separator = _guess_separator(path, head)
            return pd.read_csv(
                path,
                sep=separator,
                encoding=encoding,
                dtype=str,
                keep_default_na=False,
                na_values=[""],
                skipinitialspace=True,
            )
        except UnicodeDecodeError as error:
            last_error = error
        except pd.errors.ParserError as error:
            raise UserFacingError(
                "CSV ファイルの形が崩れていて読み込めません。",
                f"行によって列の数が違うなど、区切り文字の位置がそろっていません（{error}）。",
                "Excel で開いて、見出しが1行目にあり、各行の列数がそろっているか確かめてから保存し直してください。",
            ) from error
    raise UserFacingError(
        "CSV ファイルの文字コードを判別できません。",
        f"UTF-8 と Shift_JIS のどちらでも読めませんでした（{last_error}）。",
        "Excel で開いて「CSV UTF-8（コンマ区切り）」で保存し直してください。",
    )


def _guess_separator(path: Path, head: str) -> str:
    if path.suffix.lower() == ".tsv":
        return "\t"
    first_line = head.splitlines()[0] if head else ""
    try:
        return csv.Sniffer().sniff(first_line, delimiters=",\t;").delimiter
    except csv.Error:
        return "\t" if first_line.count("\t") > first_line.count(",") else ","


def _read_excel(path: Path, sheet: str | None, progress, cancel) -> pd.DataFrame:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise _open_error(path, error) from error
    try:
        if sheet is None:
            sheet = workbook.sheetnames[0]
        if sheet not in workbook.sheetnames:
            raise UserFacingError(
                f"シート「{sheet}」が見つかりません。",
                "ファイルを選び直したあとに、前のファイルのシート名が残っている可能性があります。",
                "シートの一覧からもう一度選んでください。",
            )
        worksheet = workbook[sheet]
        total = worksheet.max_row or 0
        header = None
        rows = []
        for count, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            if header is None:
                if row is None or all(value is None for value in row):
                    continue
                header = _unique_header(row)
                continue
            row = tuple(row[: len(header)])
            rows.append(row + (None,) * (len(header) - len(row)))
            if count % EXCEL_PROGRESS_ROWS == 0:
                check_cancel(cancel)
                if progress is not None and total:
                    progress(min(count / total, 1.0), f"Excel を読み込み中（{count:,} / {total:,} 行）")
        if header is None:
            return pd.DataFrame()
        return pd.DataFrame(rows, columns=header, dtype=object)
    finally:
        workbook.close()


def _unique_header(row) -> list[str]:
    """見出しを文字列にし、空欄と重複に名前を付ける（同じ名前の列を区別するため）。"""
    names = []
    for index, value in enumerate(row, start=1):
        name = str(value).strip() if value is not None and str(value).strip() else f"列{index}"
        base, suffix = name, 2
        while name in names:
            name = f"{base}.{suffix}"
            suffix += 1
        names.append(name)
    return names


def _open_error(path: Path, error: Exception) -> UserFacingError:
    return UserFacingError(
        "Excel ファイルを開けません。",
        f"ファイルが壊れている、パスワードで保護されている、または別のアプリが使用中の可能性があります（{type(error).__name__}）。",
        "Excel で開けるか確かめ、パスワードを外すか CSV で保存し直してから読み込んでください。",
    )
