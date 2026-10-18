"""結果の保存：Excel（シートごとの表）、CSV（UTF-8・BOM 付き。Excel でそのまま開ける）、図（PNG）、設定（JSON）。

保存先には「検証結果_年月日-時分秒」のフォルダを新しく作り、既存のファイルを上書きしない。
Excel の値はアプリで計算した結果（数式ではない）。警報は色だけでなく「警報」の文字でも示す。
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure

from . import plots
from .findings import key_findings
from .settings import method_label, save_settings

APP_NAME = "Wafer 面内傾向の異常検知 実データ検証"
APP_VERSION = "1.0.0"
ALARM_TEXT = "警報"
ALARM_FILL = "FFCDD2"
HEADER_FILL = "DDE7F0"
MAX_COLUMN_WIDTH = 60
FIGURE_DPI = 150
CHART_SIZE_INCH = (11, 4.5)
TALL_SIZE_INCH = (11, 8.5)
DATETIME_EXCEL_FORMAT = "yyyy/mm/dd hh:mm:ss"


def judgement_table(result) -> pd.DataFrame:
    """計測ごとの判定（1 Wafer の手法）。"""
    used = result.dataset.used.reset_index(drop=True)
    table = pd.DataFrame(
        {
            "順番": used["order"],
            "Wafer": used["Wafer"],
            "DateTime": used["DateTime"],
            "区分": used["segment"],
            "警報の数（1 Wafer）": result.alarms[result.single_names].sum(axis=1).to_numpy(),
            "ラベルで異常": np.where(used["label"], "異常", ""),
            "補った値の数": used["n_imputed"],
            "M3S X（nm）": result.m3s_values["M3S_X_nm"],
            "M3S Y（nm）": result.m3s_values["M3S_Y_nm"],
        }
    )
    for name in result.single_names:
        table[f"{method_label(name)} の判定"] = np.where(result.alarms[name], ALARM_TEXT, "")
        table[f"{method_label(name)} のスコア"] = result.scores[name]
    return table


def temporal_judgement_table(result) -> pd.DataFrame | None:
    if result.temporal_stats is None:
        return None
    used = result.dataset.used.reset_index(drop=True)
    table = pd.DataFrame({"順番": used["order"], "Wafer": used["Wafer"], "DateTime": used["DateTime"], "区分": used["segment"]})
    for detector in result.temporal_detectors:
        label = result.label_of(detector.name)
        stats = result.temporal_stats[detector.name]
        table[f"{label} の判定"] = np.where(result.temporal_alarms[detector.name], ALARM_TEXT, np.where(stats.isna(), "（対象外）", ""))
        table[f"{label} の統計量"] = stats
    return table


def summary_table(result, source_name: str) -> pd.DataFrame:
    dataset = result.dataset
    counts = dataset.measurements["segment"].value_counts()
    rows = [
        ("アプリ", f"{APP_NAME} {APP_VERSION}"),
        ("保存した日時", datetime.now().strftime("%Y/%m/%d %H:%M:%S")),
        ("読み込んだファイル", source_name),
        ("計測の件数", len(dataset.measurements)),
        ("共通の計測点", f"{dataset.n_points} 点（Shot {len(np.unique(dataset.points.shot_index))} 個）"),
    ]
    rows += [(f"区分：{segment}", int(count)) for segment, count in counts.items()]
    rows.append(("誤報率（1 Wafer）", f"{result.settings.methods.alpha_percent:g}%"))
    rows += [("要点", line) for line in key_findings(result)]
    if result.settings.temporal.enabled:
        rows.append(("ARL₀ の目標（経時変化）", f"{result.settings.temporal.target_arl0:g} 枚"))
    rows += [(f"計算時間：{key}", f"{value:.1f} 秒") for key, value in result.elapsed_seconds.items()]
    rows += [("注意", text) for text in result.warnings + dataset.prepared.check.warnings]
    rows += [("手法を使わなかった理由", f"{method_label(name)}：{reason}") for name, reason in result.skipped_methods]
    return pd.DataFrame(rows, columns=["項目", "内容"])


METHOD_COLUMNS = {
    "label": "手法",
    "threshold": "管理限界",
    "train_alarm_rate": "学習の警報率（参考）",
    "calibration_n": "校正の件数",
    "calibration_alarms": "校正の警報数",
    "calibration_alarm_rate": "校正の警報率",
    "evaluation_n": "評価の件数",
    "evaluation_alarms": "評価の警報数",
    "evaluation_alarm_rate": "評価の警報率",
    "evaluation_alarms_without_m3s": "評価の警報で M3S は警報なし",
    "details": "学習した中身",
    "method": "手法の名前（検証と同じ）",
}
TEMPORAL_COLUMNS = {
    "label": "検出器",
    "threshold": "管理限界",
    "bootstrap_arl0": "ARL₀（再標本化での推定、枚）",
    "limit_note": "管理限界の注意",
    "calibration_n": "校正の件数",
    "calibration_alarms": "校正の警報数",
    "expected_calibration_alarms": "校正の警報数の目安",
    "evaluation_n": "評価の件数",
    "evaluation_alarms": "評価の警報数",
    "first_evaluation_alarm_order": "評価で最初の警報（順番）",
    "first_evaluation_alarm_wafer": "評価で最初の警報（Wafer）",
    "detector": "検出器の名前（検証と同じ）",
}
LABEL_COLUMNS = {
    "label": "手法",
    "kind": "種類",
    "n_positive": "ラベルで異常の件数",
    "detected": "そのうち警報",
    "detection_rate": "検出率",
    "detection_ci_low": "検出率の95%信頼区間（下）",
    "detection_ci_high": "検出率の95%信頼区間（上）",
    "n_negative": "評価期間のラベルなしの件数",
    "false_alarms": "そのうち警報",
    "false_alarm_rate": "ラベルなしでの警報率",
    "roc_auc": "ROC-AUC",
    "auc_ci_low": "AUC の95%信頼区間（下）",
    "auc_ci_high": "AUC の95%信頼区間（上）",
    "method": "手法の名前（検証と同じ）",
}
INJECTION_SINGLE_COLUMNS = {
    "pattern": "パターン",
    "pattern_label": "内容",
    "amplitude_sigma": "大きさ（σ_dir）",
    "rms_nm": "大きさ（計測点での RMS、nm）",
    "sigma_dir_nm": "σ_dir（nm、平均）",
    "method_label": "手法",
    "n": "回数",
    "detected": "検出した回数",
    "detection_rate": "検出率",
    "ci_low": "95%信頼区間（下）",
    "ci_high": "95%信頼区間（上）",
    "method": "手法の名前（検証と同じ）",
}
INJECTION_SINGLE_SUMMARY_COLUMNS = {
    "method_label": "手法",
    "mean_detection": "平均の検出率",
    "false_alarm_rate": "注入なしでの警報率",
    "mean_detection_m3s_blind": "M3S が10%未満の条件での平均検出率",
    "n_m3s_blind_conditions": "その条件の数",
    "method": "手法の名前（検証と同じ）",
}
INJECTION_TEMPORAL_COLUMNS = {
    "component_label": "成分",
    "profile_label": "変化の形",
    "amplitude_sigma": "大きさ（σ_dir）",
    "sigma_dir_nm": "σ_dir（nm）",
    "detector_label": "検出器",
    "n_series": "系列数",
    "n_clean": "変化前に警報がなかった系列",
    "pre_change_false_alarm_rate": "変化前に警報が出た割合",
    "detection_rate": "検出率（観測した枚数の中）",
    "detection_ci_low": "検出率の95%信頼区間（下）",
    "detection_ci_high": "検出率の95%信頼区間（上）",
    "rmst_delay": "平均の検出遅延 RMST（枚）",
    "median_delay": "遅延の中央値（枚）",
    "window": "観測した枚数",
    "detector": "検出器の名前（検証と同じ）",
}
INJECTION_TEMPORAL_SUMMARY_COLUMNS = {
    "detector_label": "検出器",
    "mean_rmst_delay": "平均の検出遅延 RMST（枚）",
    "mean_detection_rate": "平均の検出率",
    "pre_change_false_alarm_rate": "変化前に警報が出た割合",
    "null_rmst_delay": "変化なしでの RMST（参考）",
    "detector": "検出器の名前（検証と同じ）",
}


def _rename(table: pd.DataFrame | None, columns: dict) -> pd.DataFrame | None:
    if table is None or table.empty:
        return None
    present = [column for column in columns if column in table.columns]
    return table[present].rename(columns=columns)


def result_tables(result, source_name: str = "") -> dict[str, pd.DataFrame]:
    """シート名 → 表（保存する順番）。"""
    dataset = result.dataset
    check = dataset.prepared.check
    tables = {
        "概要": summary_table(result, source_name),
        "データ確認": pd.DataFrame({"内容": check.summary_lines() + [f"注意：{text}" for text in check.warnings]}),
        "手法と管理限界": _rename(result.method_table, METHOD_COLUMNS),
        "判定_1Wafer": judgement_table(result),
    }
    temporal = temporal_judgement_table(result)
    if temporal is not None:
        tables["経時変化の管理限界"] = _rename(result.temporal_table, TEMPORAL_COLUMNS)
        tables["判定_経時変化"] = temporal
    if result.label_table is not None:
        tables["ラベルとの照合"] = _rename(result.label_table, LABEL_COLUMNS)
    if result.injection_single is not None:
        single = result.injection_single.assign(method_label=result.injection_single["method"].map(method_label))
        summary = result.injection_single_summary.assign(method_label=result.injection_single_summary["method"].map(method_label))
        tables["注入試験_1Wafer_まとめ"] = _rename(summary, INJECTION_SINGLE_SUMMARY_COLUMNS)
        tables["注入試験_1Wafer"] = _rename(single, INJECTION_SINGLE_COLUMNS)
    if result.injection_temporal is not None:
        tables["注入試験_経時変化_まとめ"] = _rename(result.injection_temporal_summary, INJECTION_TEMPORAL_SUMMARY_COLUMNS)
        tables["注入試験_経時変化"] = _rename(result.injection_temporal, INJECTION_TEMPORAL_COLUMNS)
    excluded = dataset.measurements[~dataset.measurements["used"]]
    if not excluded.empty:
        tables["除外した計測"] = excluded[["order", "Wafer", "DateTime", "missing_percent", "reason"]].rename(
            columns={"order": "順番", "missing_percent": "欠けの割合（%）", "reason": "理由"}
        )
    tables["共通の計測点"] = dataset.map_points.rename(columns={"presence_percent": "正常期間で計測できた割合（%）"})
    return {name: table for name, table in tables.items() if table is not None}


def save_figure(figure: Figure, path: Path) -> None:
    FigureCanvasAgg(figure)
    figure.savefig(path, dpi=FIGURE_DPI)


def export_results(result, parent_folder: str | Path, source_name: str = "", progress=None) -> Path:
    folder = Path(parent_folder) / f"検証結果_{datetime.now():%Y%m%d-%H%M%S}"
    folder.mkdir(parents=True, exist_ok=False)
    figures = folder / "図"
    figures.mkdir()
    tables = result_tables(result, source_name)
    if progress is not None:
        progress(0.05, "Excel ファイルを書いています")
    _write_excel(tables, folder / "検証結果.xlsx")
    csv_folder = folder / "CSV"
    csv_folder.mkdir()
    for name, table in tables.items():
        table.to_csv(csv_folder / f"{name}.csv", index=False, encoding="utf-8-sig", date_format="%Y/%m/%d %H:%M:%S")
    save_settings(result.settings, folder / "設定.json")
    plots.configure_fonts()
    names = list(result.single_names) + [detector.name for detector in result.temporal_detectors]
    for number, name in enumerate(names, start=1):
        if progress is not None:
            progress(0.2 + 0.7 * number / len(names), f"図を保存しています（{number} / {len(names)}）")
        figure = Figure(figsize=CHART_SIZE_INCH)
        plots.control_chart(figure, result, name)
        save_figure(figure, figures / f"管理図_{number:02d}_{_safe(result.label_of(name))}.png")
    figure = Figure(figsize=TALL_SIZE_INCH)
    plots.alarm_overview(figure, result)
    save_figure(figure, figures / "警報の一覧.png")
    if result.injection_single is not None:
        figure = Figure(figsize=TALL_SIZE_INCH)
        plots.injection_single_chart(figure, result.injection_single)
        save_figure(figure, figures / "注入試験_1Wafer.png")
    if result.injection_temporal_summary is not None and not result.injection_temporal_summary.empty:
        figure = Figure(figsize=TALL_SIZE_INCH)
        plots.injection_temporal_chart(figure, result.injection_temporal_summary)
        save_figure(figure, figures / "注入試験_経時変化.png")
    (folder / "このフォルダについて.txt").write_text(_readme_text(source_name), encoding="utf-8")
    if progress is not None:
        progress(1.0, "保存しました")
    return folder


def _safe(text: str) -> str:
    for character in '\\/:*?"<>|｜ 　':
        text = text.replace(character, "_")
    return text[:80]


def _write_excel(tables: dict[str, pd.DataFrame], path: Path) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, table in tables.items():
            table.to_excel(writer, sheet_name=name[:31], index=False)
            sheet = writer.sheets[name[:31]]
            sheet.freeze_panes = "A2"
            header_fill = PatternFill("solid", fgColor=HEADER_FILL)
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.fill = header_fill
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            alarm_fill = PatternFill("solid", fgColor=ALARM_FILL)
            for index, column in enumerate(table.columns, start=1):
                letter = get_column_letter(index)
                values = [str(value) for value in table[column].head(200)]
                width = min(MAX_COLUMN_WIDTH, max(8, int(max([len(str(column)) * 0.6] + [len(v) * 1.1 for v in values]))))
                sheet.column_dimensions[letter].width = width
                if pd.api.types.is_datetime64_any_dtype(table[column]):
                    for cell in sheet[letter][1:]:
                        cell.number_format = DATETIME_EXCEL_FORMAT
                if str(column).endswith("の判定"):
                    for cell in sheet[letter][1:]:
                        if cell.value == ALARM_TEXT:
                            cell.fill = alarm_fill


def _readme_text(source_name: str) -> str:
    return "\n".join(
        [
            f"{APP_NAME} {APP_VERSION} の結果です。",
            f"読み込んだファイル：{source_name}",
            "",
            "検証結果.xlsx … すべての表（シートごと）。値はアプリで計算した結果で、数式ではありません。",
            "CSV/ … 同じ表の CSV（UTF-8・BOM 付き）。",
            "図/ … 管理図（手法ごと）、警報の一覧、注入試験の図（PNG）。",
            "設定.json … このときの設定。アプリの「ファイル → 設定を読み込む」で同じ設定に戻せます。",
            "",
            "実データの結果を含むので、社外のサービスに送ったり、公開の場所に置いたりしないでください。",
        ]
    )
