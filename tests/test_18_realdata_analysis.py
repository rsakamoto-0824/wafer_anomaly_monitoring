"""実データ検証アプリ：同じ誤報率での判定、OR 統合、経時変化の管理限界、注入試験、保存（練習用の模擬データで）。"""
from __future__ import annotations

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from src.realdata.analysis import run_analysis
from src.realdata.dataset import SEGMENT_CALIBRATION
from src.realdata.errors import AnalysisCancelled, UserFacingError
from src.realdata.export import export_results
from src.realdata.gui.app import quick_settings
from src.realdata.methods import OrRule
from src.realdata.preprocess import prepare
from src.realdata.sample_data import SINGLE_EVENTS, generate_sample_table
from src.realdata.settings import AppSettings, load_settings, save_settings, validate_settings
from src.realdata.temporal import block_bootstrap_indices

N_LOTS = 12  # 300 枚：正常期間 150（学習 75・校正 75）


@pytest.fixture(scope="module")
def sample():
    table, events = generate_sample_table(N_LOTS)
    return table, events


@pytest.fixture(scope="module")
def result(sample):
    table, _ = sample
    settings = quick_settings(table)
    prepared = prepare(table, settings.data)
    return run_analysis(prepared, settings)


def test_sample_has_standard_columns_and_missing_marks(sample):
    table, events = sample
    assert list(table.columns[:14]) == ["Wafer", "DateTime", "shotNo", "Direction", "Xmm", "Ymm", "Radius", "markNo",
                                        "markXmm", "markYmm", "measX", "measY", "StsX", "StsY"]
    assert ((table["StsX"] == -1) & (table["measX"] == 0)).any()
    assert len(events) == len(SINGLE_EVENTS) + 1


def test_calibration_alarm_rate_is_at_most_alpha(result):
    calibration = result.dataset.rows(SEGMENT_CALIBRATION)
    for name in result.single_names:
        rate = result.alarms[name].to_numpy()[calibration].mean()
        assert rate <= result.alpha + 1e-12, name


def test_or_rule_is_any_member_over_its_limit():
    rng = np.random.default_rng(3)
    calibration = rng.standard_normal((200, 3))
    rule = OrRule(["a", "b", "c"]).fit(calibration, 0.05)
    assert np.mean(rule.transform(calibration) > 0) <= 0.05
    new = rng.standard_normal((1000, 3))
    flags = rule.transform(new) > 0
    assert np.array_equal(flags, (new > rule.limits_).any(axis=1))
    big = np.array([[100.0, 0.0, 0.0]])  # 校正の最大値を超える値でも警報になる（経験分布関数の頭打ちがない）
    assert rule.transform(big)[0] > 0


def test_known_single_wafer_anomalies_are_detected(result):
    labels = result.label_table.set_index("method")
    assert labels["n_positive"].iloc[0] == len(SINGLE_EVENTS)
    best = labels[labels["kind"] == "1 Wafer"]["detected"].max()
    assert best >= 3  # 8σ の局所・単一 Shot などは、どれかの手法で見つかる


def test_temporal_limits_and_reset(result):
    table = result.temporal_table
    assert len(table) == len(result.temporal_detectors) > 0
    assert np.all(np.isfinite(result.temporal_thresholds))
    calibration = result.dataset.rows(SEGMENT_CALIBRATION)
    stats = result.temporal_stats.to_numpy()
    assert np.isnan(stats[: calibration.min()]).all()  # 学習期間は対象外
    assert np.isfinite(stats[calibration.min():]).all()
    # 警報のあとは最初からやり直すので、校正期間の警報は件数 ÷ ARL₀ の数倍に収まる
    assert (table["calibration_alarms"] <= 5 + 5 * table["expected_calibration_alarms"]).all()


def test_injection_detection_increases_with_amplitude(result):
    table = result.injection_single
    m3s = table[(table["pattern"] == "B1") & (table["method"] == "M3S_XY")].sort_values("amplitude_sigma")
    assert m3s["detection_rate"].iloc[-1] > m3s["detection_rate"].iloc[0]
    b10 = table[(table["pattern"] == "B10") & (table["method"] == "M3S_XY")]
    normal = table[(table["pattern"] == "正常（注入なし）") & (table["method"] == "M3S_XY")]
    assert b10["detection_rate"].iloc[0] == pytest.approx(normal["detection_rate"].iloc[0])  # 平均・σ が同じなら M3S は同じ判定
    drift = result.injection_temporal
    assert set(drift["component"]) == {"none", "low", "local"}
    assert drift["rmst_delay"].between(0, drift["window"].max()).all()


def test_block_bootstrap_keeps_consecutive_wafers():
    index = block_bootstrap_indices(50, 4, 60, 10, np.random.default_rng(0))
    assert index.shape == (4, 60) and index.min() >= 0 and index.max() < 50
    assert np.all(np.diff(index[:, :10], axis=1) == 1)


def test_export_writes_excel_csv_figures_and_settings(result, tmp_path):
    folder = export_results(result, tmp_path, "sample.csv")
    workbook = load_workbook(folder / "検証結果.xlsx", read_only=True)
    assert {"概要", "データ確認", "手法と管理限界", "判定_1Wafer", "経時変化の管理限界", "注入試験_1Wafer"} <= set(workbook.sheetnames)
    workbook.close()
    judgement = pd.read_csv(folder / "CSV" / "判定_1Wafer.csv", encoding="utf-8-sig")
    assert len(judgement) == len(result.dataset.used)
    assert (folder / "図" / "警報の一覧.png").stat().st_size > 10_000
    assert load_settings(folder / "設定.json").to_dict() == result.settings.to_dict()


def test_settings_round_trip_and_validation(tmp_path):
    settings = AppSettings()
    settings.methods.or_members = ["M3S_XY", "MD_d3"]  # 選んでいない手法
    problems = validate_settings(settings)
    assert any("OR 統合" in problem for problem in problems)
    save_settings(settings, tmp_path / "s.json")
    assert load_settings(tmp_path / "s.json").to_dict() == settings.to_dict()
    assert validate_settings(AppSettings()) == []


def test_cancel_stops_analysis(sample):
    import threading

    table, _ = sample
    settings = quick_settings(table)
    prepared = prepare(table, settings.data)
    cancel = threading.Event()
    cancel.set()
    with pytest.raises(AnalysisCancelled):
        run_analysis(prepared, settings, cancel=cancel)


def test_invalid_settings_are_reported_in_plain_words(sample):
    table, _ = sample
    settings = quick_settings(table)
    settings.split.train_fraction_percent = 95
    with pytest.raises(UserFacingError) as error:
        run_analysis(prepare(table, settings.data), settings)
    assert "学習に使う割合" in error.value.why
